import os
import fitz  # PyMuPDF
import docx
import logging

logger = logging.getLogger(__name__)

def extract_text_from_pdf(file_path: str) -> list[dict]:
    """
    Extracts text page-by-page from a PDF file.
    Returns a list of dicts containing page number and text.
    """
    try:
        doc = fitz.open(file_path)
        pages_content = []
        for page_num in range(len(doc)):
            page = doc.load_page(page_num)
            text = page.get_text()
            if text.strip():
                pages_content.append({
                    "page": page_num + 1,
                    "text": text
                })
        doc.close()
        return pages_content
    except Exception as e:
        logger.warning(f"Error extracting PDF text from {file_path}: {e}. Trying fallback plain text extraction.")
        try:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                content = f.read()
                if content.strip():
                    return [{"page": 1, "text": content}]
        except Exception as fallback_e:
            logger.error(f"Fallback plain text extraction failed: {fallback_e}")
        raise e

def extract_text_from_docx(file_path: str) -> list[dict]:
    """
    Extracts text from a DOCX file.
    Returns a list of dicts with page number (always 1 for docx) and text.
    """
    try:
        doc = docx.Document(file_path)
        paragraphs_text = []
        
        # Extract paragraph texts
        for para in doc.paragraphs:
            if para.text.strip():
                paragraphs_text.append(para.text)
                
        # Extract table texts
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    paragraphs_text.append(row_text)
                    
        full_text = "\n\n".join(paragraphs_text)
        return [{"page": 1, "text": full_text}]
    except Exception as e:
        logger.warning(f"Error extracting DOCX text from {file_path}: {e}. Trying fallback plain text extraction.")
        try:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                content = f.read()
                if content.strip():
                    return [{"page": 1, "text": content}]
        except Exception as fallback_e:
            logger.error(f"Fallback plain text extraction failed: {fallback_e}")
        raise e

def extract_text(file_path: str) -> list[dict]:
    """
    Main entry point for document text extraction.
    """
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"File not found: {file_path}")
        
    ext = os.path.splitext(file_path)[1].lower()
    if ext == '.pdf':
        return extract_text_from_pdf(file_path)
    elif ext == '.docx':
        return extract_text_from_docx(file_path)
    else:
        raise ValueError(f"Unsupported file type: {ext}")
